import os
import io
import requests
import hashlib
import urllib3
import re
import logging
import openpyxl
from urllib.parse import urljoin
from datetime import datetime
from django.utils import timezone
from django.core.files.base import ContentFile
from apps.documents.models import RuleSourceDocument, OECDSyncLog
from apps.rule_extraction.tasks import process_document_task

urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)
logger = logging.getLogger(__name__)

def parse_date(date_str):
    try:
        if isinstance(date_str, datetime):
            return date_str
        return datetime.strptime(str(date_str).strip(), "%d %b '%y %H:%M")
    except Exception:
        return timezone.now()

def scan_oecd_from_excel():
    update_sheet_url = "https://www.oecd.org/content/dam/oecd/en/topics/policy-issue-focus/aeoi/aeoi-portal-updates.xlsx"
    headers = {'User-Agent': 'Mozilla/5.0'}
    
    resp = requests.get(update_sheet_url, headers=headers, verify=False, timeout=20)
    resp.raise_for_status()
    
    wb = openpyxl.load_workbook(io.BytesIO(resp.content), data_only=True)
    sheet = wb.active
    
    found_pdfs = {}
    
    for row in sheet.iter_rows(values_only=True):
        if not row or not row[0]: continue
        content_val = str(row[0])
        
        if 'tin' in content_val.lower():
            country_raw = content_val.lower().replace('- tin', '').replace(' tin', '').strip()
            slug = country_raw.replace(' ', '-')
            
            if slug == 'unied-kingdom': slug = 'united-kingdom'
            if slug == 'unied-arab-emiraes': slug = 'united-arab-emirates'
            if slug == 'swizerland': slug = 'switzerland'
            if slug == 'mauriius': slug = 'mauritius'
            if slug == 'gibralar': slug = 'gibraltar'
            if slug == 'monserra': slug = 'montserrat'
            if slug == 'lihuania': slug = 'lithuania'
            
            pdf_url = f"https://www.oecd.org/content/dam/oecd/en/topics/policy-issue-focus/aeoi/{slug}-tin.pdf"
            found_pdfs[slug] = pdf_url
            
    if not found_pdfs:
        raise Exception("Failed to parse any TIN entries from Excel.")
        
    return found_pdfs

def run_oecd_sync(log_id: int, user_id=None):
    try:
        log = OECDSyncLog.objects.get(id=log_id)
    except OECDSyncLog.DoesNotExist:
        return

    from apps.users.models import User
    user = User.objects.filter(id=user_id).first() if user_id else None

    try:
        found_pdfs = scan_oecd_from_excel()
    except Exception as e:
        log.status = 'failed'
        log.error_details = {'network_error': str(e)}
        log.end_time = timezone.now()
        log.save()
        return

    log.total_found = len(found_pdfs)
    log.save(update_fields=['total_found'])

    error_details = {}
    downloaded_count = 0
    error_count = 0
    
    headers = {'User-Agent': 'Mozilla/5.0'}

    for country, pdf_url in found_pdfs.items():
        try:
            resp = requests.get(pdf_url, headers=headers, timeout=15, verify=False)
            if resp.status_code == 200:
                print(f"Downloaded {country}")
                content = resp.content
                file_hash = hashlib.md5(content).hexdigest()
                
                doc = RuleSourceDocument.objects.filter(country__iexact=country).first()
                if doc:
                    if doc.file_hash == file_hash:
                        print(f"  Unchanged: {country}")
                        continue
                    else:
                        print(f"  Updated: {country}")
                        is_new = False
                else:
                    print(f"  New: {country}")
                    is_new = True
                    doc = RuleSourceDocument(
                        title=f"{country.title()} TIN Rules",
                        country=country.title(),
                        source_url=pdf_url,
                        document_type='oecd_pdf',
                        uploaded_by=user
                    )
                
                doc.file_hash = file_hash
                doc.is_processed = False
                
                filename = f"{country}_tin.pdf"
                if doc.file:
                    doc.file.delete(save=False)
                doc.file.save(filename, ContentFile(content), save=False)
                doc.save()
                
                process_document_task.delay(doc.id)
                downloaded_count += 1
            else:
                error_count += 1
        except Exception as e:
            error_count += 1
            error_details[country] = str(e)

    log.downloaded_count = downloaded_count
    log.error_count = error_count
    log.error_details = error_details
    log.end_time = timezone.now()
    log.status = 'completed' if error_count == 0 else 'partial'
    log.save()
